import os
import json
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document


# ============================================================
# GROQ CONFIGURATION
# ============================================================

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError(
        "GROQ_API_KEY is missing. Please add it to your .env file."
    )

client = Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"


# ============================================================
# JOB DESCRIPTION MODEL
# ============================================================

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


# ============================================================
# RESUME MODELS
# ============================================================

class MatchResult(BaseModel):
    score: float
    details: dict


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


# ============================================================
# EXTRACT JOB INFORMATION
# ============================================================

def extract_job_information(job_description):

    job_schema = JobD.model_json_schema()

    system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze a job description and extract
structured information from it.

Return ONLY valid JSON matching this schema:

{job_schema}

IMPORTANT:

- Do NOT return the schema itself.
- Do NOT return fields like properties, title or type.
- Fill the schema with actual information extracted
  from the job description.
- If minimum experience is not mentioned, return null.
- If information for a list is missing, return an empty list.
- Do not invent information.
"""

    user_prompt = f"""
Analyze the following job description:

{job_description}
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={
            "type": "json_object"
        }
    )

    raw_json = response.choices[0].message.content

    job_data = json.loads(raw_json)

    return JobD(**job_data)


# ============================================================
# READ PDF
# ============================================================

def read_pdf(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ============================================================
# READ DOCX
# ============================================================

def read_docx(file_path):

    document = Document(file_path)

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# ============================================================
# READ RESUME
# ============================================================

def read_resume(file_path):

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    if extension == ".pdf":

        return read_pdf(file_path)

    elif extension == ".docx":

        return read_docx(file_path)

    else:

        return None


# ============================================================
# PARSE RESUME USING GROQ
# ============================================================

def parse_resume(resume_text):

    resume_schema = Resume.model_json_schema()

    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning,
not only based on exact section headings.

Different resumes may use different headings.

For example:

- Experience
- Professional Experience
- Work History
- Employment
- Internships

These may all contain relevant experience.

Skills may also appear in:

- Skills section
- Work experience
- Internships
- Projects

Return ONLY valid JSON matching this schema:

{resume_schema}

Important rules:

1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
"""

    user_prompt = f"""
Parse the following resume:

{resume_text}
"""

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={
            "type": "json_object"
        }
    )

    raw_output = response.choices[0].message.content

    data = json.loads(raw_output)

    return Resume(**data)


# ============================================================
# COMPARE RESUME WITH JOB
# ============================================================

def final_score(job, resume):

    match_schema = MatchResult.model_json_schema()

    prompt = f"""
You are an HR recruiter and ATS resume evaluator.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:

{job.model_dump_json(indent=2)}


CANDIDATE RESUME:

{resume.model_dump_json(indent=2)}


Return JSON matching this schema:

{match_schema}


Give the following information:

1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether the experience requirement is met
5. Overall match percentage from 0 to 100
6. A short final verdict
7. Suggestions for improving the resume for this job

Keep the response concise and easy to understand.

The score should be based only on the information present
in the job description and candidate resume.
Do not invent candidate information.
"""

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        response_format={
            "type": "json_object"
        }
    )

    data = json.loads(
        response.choices[0].message.content
    )

    return MatchResult(**data)


# ============================================================
# OPTIONAL: RUN AS NORMAL PYTHON SCRIPT
# ============================================================

if __name__ == "__main__":

    print("Resume Parser backend is working.")

    print(
        "Use app.py to run the Streamlit web application."
    )
